from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document("/root/arahin/laporan/laporan_metopen_v3.docx")

# Set default font and spacing
style = doc.styles['Normal']
font = style.font
font.name = 'Times New Roman'
font.size = Pt(12)
font.color.rgb = RGBColor(0, 0, 0)

# Set paragraph spacing
para_format = style.paragraph_format
para_format.line_spacing = 1.5
para_format.space_after = Pt(6)
para_format.space_before = Pt(0)

# Set margins
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3)
    section.right_margin = Cm(3)

# Format all paragraphs
for para in doc.paragraphs:
    text = para.text.strip()
    style_name = para.style.name
    
    # Set font for each run
    for run in para.runs:
        run.font.name = 'Times New Roman'
    
    # Format based on style
    if style_name == 'Title':
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(16)
            run.font.bold = True
    
    elif style_name == 'Author':
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in para.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
    
    elif style_name == 'Heading 1':
        for run in para.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(14)
            run.font.bold = True
        para.paragraph_format.space_before = Pt(18)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.5
    
    elif style_name == 'Heading 2':
        for run in para.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.5
    
    elif style_name == 'Heading 3':
        for run in para.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
            run.font.bold = True
        para.paragraph_format.space_before = Pt(12)
        para.paragraph_format.space_after = Pt(6)
        para.paragraph_format.line_spacing = 1.5
    
    elif style_name in ['Body Text', 'First Paragraph', 'Abstract']:
        for run in para.runs:
            run.font.name = 'Times New Roman'
            run.font.size = Pt(12)
        para.paragraph_format.line_spacing = 1.5

# Format tables
for table in doc.tables:
    for row in table.rows:
        for cell in row.cells:
            cell_para = cell.paragraphs[0]
            for run in cell_para.runs:
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
            cell_para.paragraph_format.line_spacing = 1.0
            cell_para.paragraph_format.space_before = Pt(2)
            cell_para.paragraph_format.space_after = Pt(2)

doc.save("/root/arahin/laporan/laporan_metopen_v3_formatted.docx")
print("Done!")
